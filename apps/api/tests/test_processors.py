import json
import zipfile

import pytest

from pluto.parsers import checked_office, extract


def test_streaming_text_beyond_previous_limit(tmp_path):
    path = tmp_path / "large.txt"
    with path.open("wb") as output:
        for _ in range(17000):
            output.write(b"x" * 4095 + b"\n")
    assert sum(1 for _ in extract(path, "large.txt")) == 17000


@pytest.mark.parametrize("kind", ["csv", "json", "jsonl"])
def test_structured(tmp_path, kind):
    path = tmp_path / kind
    path.write_text(
        {"csv": "id,name\n1,Pluto\n", "json": '[{"name":"Pluto"}]', "jsonl": '{"name":"Pluto"}\n'}[
            kind
        ]
    )
    assert "Pluto" in str(list(extract(path, "data." + kind)))


def test_office_and_parquet(tmp_path):
    import polars as pl
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation

    doc = Document()
    doc.add_paragraph("Document evidence")
    p = tmp_path / "test.docx"
    doc.save(p)
    assert "Document evidence" in str(list(extract(p, p.name)))
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = "Slide evidence"
    p = tmp_path / "test.pptx"
    slides.save(p)
    assert "Slide evidence" in str(list(extract(p, p.name)))
    book = Workbook()
    book.active.append(["Sheet evidence"])
    p = tmp_path / "test.xlsx"
    book.save(p)
    assert "Sheet evidence" in str(list(extract(p, p.name)))
    p = tmp_path / "test.parquet"
    pl.DataFrame({"value": ["Parquet evidence"]}).write_parquet(p)
    assert "Parquet evidence" in str(list(extract(p, p.name)))


def test_pdf_and_image(tmp_path, monkeypatch):
    import pymupdf
    from PIL import Image

    import pluto.parsers as parsers

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((30, 30), "PDF evidence")
    p = tmp_path / "test.pdf"
    document.save(p)
    document.close()
    preview = tmp_path / "preview.png"
    assert "PDF evidence" in str(list(extract(p, p.name, preview))) and preview.exists()
    monkeypatch.setattr(parsers, "ocr", lambda image: "Local OCR test adapter")
    p = tmp_path / "test.png"
    Image.new("RGB", (20, 30)).save(p)
    content = list(extract(p, p.name, preview))
    assert json.loads(content[0][1])["width"] == 20
    assert content[1][0] == "image OCR"


def test_archive_traversal(tmp_path):
    p = tmp_path / "bad.docx"
    with zipfile.ZipFile(p, "w") as archive:
        archive.writestr("../outside", "bad")
    with pytest.raises(ValueError, match="Unsafe archive"):
        checked_office(p)
